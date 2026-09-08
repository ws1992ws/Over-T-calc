
import streamlit as st
import pandas as pd
from datetime import date, timedelta
from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Overtime Calculator", page_icon="🧮", layout="wide")
st.title("🧮 Overtime Calculator")
st.caption("Overtime calculator with effective-date salary/allowance history, exclusions, save/load, and Excel export.")

DEFAULTS = {"weekday_factor":3.5,"weekend_factor":7.0,"salary_divisor":155.0,"transport_divisor":30.0}

for k,v in DEFAULTS.items():
    st.session_state.setdefault(k,v)
st.session_state.setdefault("salary_data", pd.DataFrame({"Effective Date":[date(date.today().year,1,1)],"Basic Salary (SAR)":[0.0]}))
st.session_state.setdefault("transport_data", pd.DataFrame({"Effective Date":[date(date.today().year,1,1)],"Transportation Allowance (SAR)":[0.0]}))
st.session_state.setdefault("exclusion_data", pd.DataFrame(columns=["Type","Start Date","End Date"]))
st.session_state.setdefault("calc_result", None)

def hist(df, dc, ac):
    x=df.dropna(subset=[dc,ac]).copy()
    if x.empty: return []
    x[dc]=pd.to_datetime(x[dc]).dt.date
    x[ac]=pd.to_numeric(x[ac],errors="coerce")
    x=x.dropna(subset=[ac]).sort_values(dc).drop_duplicates(dc,keep="last")
    return list(zip(x[dc],x[ac]))

def effective(h,d):
    a=[x for x in h if x[0] <= d]
    return max(a,key=lambda x:x[0])[1] if a else None

def exclusions(df):
    if df.empty: return []
    x=df.dropna(subset=["Type","Start Date","End Date"]).copy()
    if x.empty: return []
    x["Start Date"]=pd.to_datetime(x["Start Date"]).dt.date
    x["End Date"]=pd.to_datetime(x["End Date"]).dt.date
    out=[]
    for _,r in x.iterrows():
        if r["End Date"] < r["Start Date"]:
            raise ValueError("An exclusion end date is before its start date.")
        out.append((str(r["Type"]),r["Start Date"],r["End Date"]))
    return out

def calc(start,end,sdf,tdf,edf,wf,wef,sd,td):
    if end < start: raise ValueError("End date cannot be before start date.")
    if sd == 0 or td == 0: raise ValueError("Divisors cannot be zero.")
    sh=hist(sdf,"Effective Date","Basic Salary (SAR)")
    th=hist(tdf,"Effective Date","Transportation Allowance (SAR)")
    ex=exclusions(edf)
    if not sh: raise ValueError("Enter at least one Basic Salary.")
    if not th: raise ValueError("Enter at least one Transportation Allowance.")
    rows=[]; d=start
    while d <= end:
        reason=" / ".join(dict.fromkeys([t for t,a,b in ex if a <= d <= b]))
        if reason:
            rows.append([d,d.strftime("%A"),"Excluded",reason,None,None,"Excluded",0.0,0.0,0.0])
        else:
            b=effective(sh,d); t=effective(th,d)
            if b is None: raise ValueError(f"No Basic Salary is effective on {d}.")
            if t is None: raise ValueError(f"No Transportation Allowance is effective on {d}.")
            weekend=d.weekday() in (4,5)
            typ="Weekend" if weekend else "Weekday"
            base=b*(wef if weekend else wf)/sd
            tr=t/td
            rows.append([d,d.strftime("%A"),"Included","",float(b),float(t),typ,float(base),float(tr),float(base+tr)])
        d += timedelta(days=1)
    cols=["Date","Day","Status","Exclusion Reason","Basic Salary (SAR)","Transportation Allowance (SAR)","OT Type","Base OT (SAR)","Transportation OT (SAR)","Daily Total (SAR)"]
    rdf=pd.DataFrame(rows,columns=cols)
    inc=rdf[rdf["Status"]=="Included"].copy()
    if inc.empty:
        mdf=pd.DataFrame(columns=["Month","Weekday Days","Weekend Days","Weekday OT (SAR)","Weekend OT (SAR)","Transportation OT (SAR)","Monthly Total (SAR)"])
    else:
        inc["Month"]=pd.to_datetime(inc["Date"]).dt.to_period("M").astype(str)
        rec=[]
        for m,g in inc.groupby("Month",sort=True):
            rec.append({
                "Month":m,
                "Weekday Days":int((g["OT Type"]=="Weekday").sum()),
                "Weekend Days":int((g["OT Type"]=="Weekend").sum()),
                "Weekday OT (SAR)":g.loc[g["OT Type"]=="Weekday","Base OT (SAR)"].sum(),
                "Weekend OT (SAR)":g.loc[g["OT Type"]=="Weekend","Base OT (SAR)"].sum(),
                "Transportation OT (SAR)":g["Transportation OT (SAR)"].sum(),
                "Monthly Total (SAR)":g["Daily Total (SAR)"].sum()
            })
        mdf=pd.DataFrame(rec)
    summary={
        "total":float(inc["Daily Total (SAR)"].sum()) if not inc.empty else 0.0,
        "weekday_days":int((inc["OT Type"]=="Weekday").sum()) if not inc.empty else 0,
        "weekend_days":int((inc["OT Type"]=="Weekend").sum()) if not inc.empty else 0,
        "excluded_days":int((rdf["Status"]=="Excluded").sum()),
        "weekday_ot":float(inc.loc[inc["OT Type"]=="Weekday","Base OT (SAR)"].sum()) if not inc.empty else 0.0,
        "weekend_ot":float(inc.loc[inc["OT Type"]=="Weekend","Base OT (SAR)"].sum()) if not inc.empty else 0.0,
        "transport_ot":float(inc["Transportation OT (SAR)"].sum()) if not inc.empty else 0.0
    }
    return rdf,mdf,summary

def excel_bytes(rdf,mdf,summary,start,end,rules):
    wb=Workbook(); ws=wb.active; ws.title="Summary"
    blue=PatternFill("solid",fgColor="1F4E78"); white=Font(color="FFFFFF",bold=True)
    thin=Side(style="thin",color="D0D7DE"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
    ws["A1"]="Overtime Calculation Report"; ws["A1"].font=Font(size=18,bold=True); ws.merge_cells("A1:D1")
    ws["A3"]="Period"; ws["B3"]=f"{start} to {end}"
    data=[("Total Overtime (SAR)",summary["total"]),("Weekday OT Days",summary["weekday_days"]),("Weekend OT Days",summary["weekend_days"]),("Excluded Days",summary["excluded_days"]),("Weekday OT (SAR)",summary["weekday_ot"]),("Weekend OT (SAR)",summary["weekend_ot"]),("Transportation OT (SAR)",summary["transport_ot"])]
    r=5
    for a,b in data:
        ws.cell(r,1,a).font=Font(bold=True); ws.cell(r,2,b)
        ws.cell(r,1).border=ws.cell(r,2).border=border; r+=1
    r+=1
    ws.cell(r,1,"Rules").fill=blue; ws.cell(r,1).font=white; r+=1
    rules_text=[
        ("Weekday","Sunday–Thursday"),
        ("Weekend","Friday–Saturday"),
        ("Weekday formula",f"Basic Salary × {rules['weekday_factor']} ÷ {rules['salary_divisor']}"),
        ("Weekend formula",f"Basic Salary × {rules['weekend_factor']} ÷ {rules['salary_divisor']}"),
        ("Transportation",f"Allowance ÷ {rules['transport_divisor']}"),
        ("Public holidays","Not automatically excluded")]
    for a,b in rules_text:
        ws.cell(r,1,a); ws.cell(r,2,b); r+=1
    for name,df in [("Monthly Summary",mdf),("Daily Detail",rdf)]:
        sh=wb.create_sheet(name)
        for c,h in enumerate(df.columns,1):
            cell=sh.cell(1,c,h); cell.fill=blue; cell.font=white; cell.alignment=Alignment(horizontal="center")
        for ri,row in enumerate(df.itertuples(index=False),2):
            for ci,val in enumerate(row,1):
                cell=sh.cell(ri,ci,val)
                if "SAR" in df.columns[ci-1]: cell.number_format='#,##0.00'
                if ci==1 and hasattr(val,"isoformat"): cell.number_format="yyyy-mm-dd"
        sh.freeze_panes="A2"
        for c in range(1,sh.max_column+1):
            sh.column_dimensions[get_column_letter(c)].width=min(32,max(12,max(len(str(sh.cell(rr,c).value or "")) for rr in range(1,sh.max_row+1))+2))
    out=BytesIO(); wb.save(out); return out.getvalue()

def project_bytes(start,end,sdf,tdf,edf,rules):
    def ds(v): return pd.to_datetime(v).date().isoformat()
    data={
        "period":{"start":start.isoformat(),"end":end.isoformat()},
        "rules":rules,
        "salary_history":[{"effective_date":ds(r["Effective Date"]),"amount":float(r["Basic Salary (SAR)"])} for _,r in sdf.dropna().iterrows()],
        "transport_history":[{"effective_date":ds(r["Effective Date"]),"amount":float(r["Transportation Allowance (SAR)"])} for _,r in tdf.dropna().iterrows()],
        "exclusions":[{"type":str(r["Type"]),"start":ds(r["Start Date"]),"end":ds(r["End Date"])} for _,r in edf.dropna(subset=["Type","Start Date","End Date"]).iterrows()]
    }
    return json.dumps(data,indent=2).encode("utf-8")

with st.sidebar:
    st.header("Project")
    up=st.file_uploader("Load saved project",type=["json"])
    if up is not None:
        try:
            data=json.load(up)
            st.session_state["loaded_period"]=(date.fromisoformat(data["period"]["start"]),date.fromisoformat(data["period"]["end"]))
            st.session_state.salary_data=pd.DataFrame([{"Effective Date":date.fromisoformat(x["effective_date"]),"Basic Salary (SAR)":x["amount"]} for x in data.get("salary_history",[])])
            st.session_state.transport_data=pd.DataFrame([{"Effective Date":date.fromisoformat(x["effective_date"]),"Transportation Allowance (SAR)":x["amount"]} for x in data.get("transport_history",[])])
            st.session_state.exclusion_data=pd.DataFrame([{"Type":x["type"],"Start Date":date.fromisoformat(x["start"]),"End Date":date.fromisoformat(x["end"])} for x in data.get("exclusions",[])])
            for k,v in data.get("rules",{}).items():
                if k in DEFAULTS: st.session_state[k]=float(v)
            st.success("Project loaded.")
        except Exception as e:
            st.error(f"Could not load project: {e}")

    st.divider()
    st.subheader("Calculation Rules")
    wf=st.number_input("Weekday factor",value=float(st.session_state.weekday_factor),step=0.1)
    wef=st.number_input("Weekend factor",value=float(st.session_state.weekend_factor),step=0.1)
    sd=st.number_input("Salary divisor",value=float(st.session_state.salary_divisor),step=1.0)
    td=st.number_input("Transportation divisor",value=float(st.session_state.transport_divisor),step=1.0)
    st.caption("Friday & Saturday are weekends. Public holidays are not automatically excluded.")

loaded=st.session_state.pop("loaded_period",None)
default_start=loaded[0] if loaded else date(date.today().year,1,1)
default_end=loaded[1] if loaded else date.today()

st.subheader("1. Overtime Period")
c1,c2=st.columns(2)
start=c1.date_input("Start date",value=default_start,key="start_v3")
end=c2.date_input("End date",value=default_end,key="end_v3")
st.caption("All dates are counted unless excluded as Vacation or Intedab.")

st.subheader("2. Basic Salary History")
sdf=st.data_editor(st.session_state.salary_data,num_rows="dynamic",hide_index=True,use_container_width=True,
    column_config={"Effective Date":st.column_config.DateColumn(format="YYYY-MM-DD"),"Basic Salary (SAR)":st.column_config.NumberColumn(min_value=0.0,format="%.2f")},key="sal_v3")
st.session_state.salary_data=sdf

st.subheader("3. Transportation Allowance History")
tdf=st.data_editor(st.session_state.transport_data,num_rows="dynamic",hide_index=True,use_container_width=True,
    column_config={"Effective Date":st.column_config.DateColumn(format="YYYY-MM-DD"),"Transportation Allowance (SAR)":st.column_config.NumberColumn(min_value=0.0,format="%.2f")},key="tr_v3")
st.session_state.transport_data=tdf

st.subheader("4. Excluded Periods")
edf=st.data_editor(st.session_state.exclusion_data,num_rows="dynamic",hide_index=True,use_container_width=True,
    column_config={"Type":st.column_config.SelectboxColumn(options=["Vacation","Intedab"]),"Start Date":st.column_config.DateColumn(format="YYYY-MM-DD"),"End Date":st.column_config.DateColumn(format="YYYY-MM-DD")},key="ex_v3")
st.session_state.exclusion_data=edf

rules={"weekday_factor":wf,"weekend_factor":wef,"salary_divisor":sd,"transport_divisor":td}

st.download_button("💾 Save Current Inputs",project_bytes(start,end,sdf,tdf,edf,rules),"overtime_project.json","application/json",use_container_width=True)

st.divider()
if st.button("Calculate Overtime",type="primary",use_container_width=True):
    try:
        rdf,mdf,summary=calc(start,end,sdf,tdf,edf,wf,wef,sd,td)
        st.session_state.calc_result={"rdf":rdf,"mdf":mdf,"summary":summary,"start":start,"end":end,"rules":rules}
        st.success("Calculation completed.")
    except Exception as e:
        st.error(str(e))

if st.session_state.calc_result:
    x=st.session_state.calc_result; rdf=x["rdf"]; mdf=x["mdf"]; s=x["summary"]
    st.subheader("Results Summary")
    a,b,c,d=st.columns(4)
    a.metric("Total Overtime",f"{s['total']:,.2f} SAR")
    b.metric("Weekday OT Days",s["weekday_days"])
    c.metric("Weekend OT Days",s["weekend_days"])
    d.metric("Excluded Days",s["excluded_days"])
    e,f,g=st.columns(3)
    e.metric("Weekday OT",f"{s['weekday_ot']:,.2f} SAR")
    f.metric("Weekend OT",f"{s['weekend_ot']:,.2f} SAR")
    g.metric("Transportation OT",f"{s['transport_ot']:,.2f} SAR")

    st.subheader("Monthly Summary")
    st.dataframe(mdf,use_container_width=True,hide_index=True)

    st.subheader("Daily Audit Trail")
    st.dataframe(rdf,use_container_width=True,hide_index=True)

    st.subheader("Downloads")
    c1,c2,c3=st.columns(3)
    c1.download_button("Download Excel Report",excel_bytes(rdf,mdf,s,x["start"],x["end"],x["rules"]),"Overtime_Calculation_Report.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
    c2.download_button("Download Daily CSV",rdf.to_csv(index=False).encode("utf-8-sig"),"overtime_daily_detail.csv","text/csv",use_container_width=True)
    c3.download_button("Download Monthly CSV",mdf.to_csv(index=False).encode("utf-8-sig"),"overtime_monthly_summary.csv","text/csv",use_container_width=True)

with st.expander("Calculation explanation"):
    st.markdown(f"""
- **Weekdays:** Sunday–Thursday → `Basic Salary × {wf} ÷ {sd}`
- **Weekends:** Friday–Saturday → `Basic Salary × {wef} ÷ {sd}`
- **Transportation:** `Transportation Allowance ÷ {td}`
- Salary and allowance are selected separately for each exact date.
- Vacation and Intedab dates are excluded.
- Public holidays are not automatically excluded.
""")

st.caption("The app does not intentionally save your data to a database. Use Save Current Inputs to keep a reusable local JSON project file.")
